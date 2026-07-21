"""Tests de los extractores plugin (DoD Fase 4): registro, límites, aislamiento."""

from __future__ import annotations

import io
import time

from normalizacion.core.config import PerillasWorker
from normalizacion.ingesta.workers.extractores import (
    ContextoExtraccion,
    ResultadoExtraccion,
    extractor_para,
    extraer,
    registrar,
)

PERILLAS = PerillasWorker()

CSV = b"id,nombre,monto\n1,ana,10.5\n2,luis,\n3,eva,9.99\n"


def _extraer(datos: bytes, tipo: str, perillas: PerillasWorker = PERILLAS) -> ResultadoExtraccion:
    return extraer(perillas, io.BytesIO(datos), tipo_real=tipo, nombre="x", tamano=len(datos))


class TestRegistro:
    def test_tipo_nuevo_es_solo_un_plugin(self) -> None:
        """DoD Fase 4: añadir un tipo = soltar un plugin, SIN tocar el núcleo."""

        @registrar("application/x-tipo-nuevo")
        def extraer_nuevo(ctx: ContextoExtraccion) -> ResultadoExtraccion:
            return ResultadoExtraccion(campos={"plugin": "nuevo"})

        r = _extraer(b"lo que sea", "application/x-tipo-nuevo")
        assert r.campos == {"plugin": "nuevo"}

    def test_prefijo_comodin(self) -> None:
        assert extractor_para("image/png") is not None  # registrado como image/*

    def test_sin_extractor_es_flag_no_error(self) -> None:
        r = _extraer(b"...", "application/x-desconocido-total")
        assert r.flags == ["sin_extractor_l1"]
        assert r.campos == {}


class TestAislamiento:
    def test_timeout_no_cuelga_al_worker(self) -> None:
        """DoD: un plugin colgado → flag extraccion_timeout y el worker SIGUE."""

        @registrar("application/x-lentisimo")
        def extraer_lento(ctx: ContextoExtraccion) -> ResultadoExtraccion:
            time.sleep(5)
            return ResultadoExtraccion()

        perillas = PerillasWorker(extractor_timeout_s=0.2)
        inicio = time.monotonic()
        r = _extraer(b"x", "application/x-lentisimo", perillas)
        assert time.monotonic() - inicio < 2.0
        assert r.flags == ["extraccion_timeout"]

    def test_crash_del_plugin_es_flag(self) -> None:
        """Un plugin que revienta NO tira el worker: flag con el tipo de error."""

        @registrar("application/x-explosivo")
        def extraer_explosivo(ctx: ContextoExtraccion) -> ResultadoExtraccion:
            raise ValueError("boom")

        r = _extraer(b"x", "application/x-explosivo")
        assert r.flags == ["extraccion_fallida:ValueError"]

    def test_pdf_falso_es_flag_no_excepcion(self) -> None:
        """El PDF basura del disco sintético no debe tumbar nada."""
        r = _extraer(b"%PDF-1.4\nbasura sin estructura\n%%EOF", "application/pdf")
        assert r.flags and r.flags[0].startswith("extraccion_fallida:")


class TestTabular:
    def test_csv_con_perfil_de_calidad(self) -> None:
        r = _extraer(CSV, "text/csv")
        assert r.campos["filas"] == 3
        assert r.campos["columnas"] == 3
        perfil = r.perfil_calidad
        assert perfil is not None
        assert 0 <= perfil["quality_score"] <= 100
        # la columna monto tiene 1 nulo de 3 filas
        assert perfil["columnas_detalle"]["monto"]["nulos_pct"] > 0

    def test_ndjson(self) -> None:
        r = _extraer(b'{"a": 1, "b": "x"}\n{"a": 2, "b": "y"}\n', "application/x-ndjson")
        assert r.campos["filas"] == 2
        assert r.perfil_calidad is not None

    def test_json_unico_expone_claves(self) -> None:
        r = _extraer(b'{"cliente": "ana", "total": 9}', "application/json")
        assert r.campos["claves_raiz"] == ["cliente", "total"]


class TestDocumentosReales:
    def test_pdf_real(self) -> None:
        from pypdf import PdfWriter

        buf = io.BytesIO()
        escritor = PdfWriter()
        escritor.add_blank_page(width=72, height=72)
        escritor.write(buf)
        r = _extraer(buf.getvalue(), "application/pdf")
        assert r.campos["paginas"] == 1
        assert not any(f.startswith("extraccion_fallida") for f in r.flags)

    def test_docx_real_extrae_texto(self) -> None:
        from docx import Document  # type: ignore[import-untyped]

        buf = io.BytesIO()
        d = Document()
        d.add_paragraph("Contrato de servicios de normalizacion masiva")
        d.save(buf)
        r = _extraer(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert r.campos["parrafos"] == 1
        assert r.texto is not None and "normalizacion masiva" in r.texto

    def test_xlsx_real_lista_hojas(self) -> None:
        from openpyxl import Workbook  # type: ignore[import-untyped]

        buf = io.BytesIO()
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Ventas"
        hoja.append(["id", "monto"])
        hoja.append([1, 99])
        libro.save(buf)
        r = _extraer(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        assert r.campos["hojas"] == ["Ventas"]
        assert r.campos["hojas_detalle"]["Ventas"]["filas"] == 2

    def test_imagen_real(self) -> None:
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (32, 16)).save(buf, format="PNG")
        r = _extraer(buf.getvalue(), "image/png")
        assert r.campos["ancho"] == 32 and r.campos["alto"] == 16


class TestLimites:
    def test_texto_truncado_con_flag(self) -> None:
        """Patrón fscrawler/Tika (⚙K11): se indexa lo parcial + flag, jamás se descarta."""
        perillas = PerillasWorker(extractor_max_chars=50)
        r = _extraer(b"una linea de texto cualquiera\n" * 100, "text/plain")
        del r
        r2 = extraer(
            perillas,
            io.BytesIO(b"una linea de texto cualquiera\n" * 100),
            tipo_real="text/plain",
            nombre="grande.txt",
            tamano=3000,
        )
        assert r2.texto is not None and len(r2.texto) <= 50
        assert "texto_truncado" in r2.flags


def _tesseract_disponible() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _png(texto: str, tamano: tuple[int, int] = (620, 200)) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", tamano, "white")
    dibujo = ImageDraw.Draw(img)
    try:
        fuente = ImageFont.load_default(size=48)
    except TypeError:  # Pillow viejo sin `size`
        fuente = ImageFont.load_default()
    dibujo.text((20, 70), texto, fill="black", font=fuente)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


class TestImagenOCR:
    def test_metadata_siempre_y_degrada_sin_tesseract(self) -> None:
        """Dimensiones/EXIF SIEMPRE; el OCR degrada con flag (nunca rompe)."""
        r = _extraer(_png("HOLA"), "image/png")
        assert r.campos["ancho"] == 620 and r.campos["alto"] == 200
        assert r.campos["formato"] == "PNG"
        if _tesseract_disponible():
            assert "ocr_ok" in r.flags or "ocr_vacio" in r.flags
        else:
            assert "ocr_no_disponible" in r.flags
            assert r.texto is None

    def test_imagen_pequena_se_salta(self) -> None:
        """Íconos/miniaturas (< ocr_min_lado) no gastan OCR."""
        r = _extraer(_png("x", tamano=(20, 20)), "image/png")
        assert "ocr_saltado_pequena" in r.flags
        assert r.texto is None

    def test_imagen_corrupta_no_rompe(self) -> None:
        """Bytes que no son imagen → flag, jamás excepción (garantía del despacho)."""
        r = _extraer(b"\x89PNG\r\n\x1a\n basura no-imagen", "image/png")
        assert any(f.startswith("extraccion_fallida") for f in r.flags)

    def test_ocr_extrae_texto(self) -> None:
        """Con Tesseract instalado, el texto de la imagen llega a `texto` (→ texto_indexable)."""
        if not _tesseract_disponible():
            import pytest

            pytest.skip("tesseract no instalado en este entorno")
        r = _extraer(_png("FACTURA 12345"), "image/png")
        assert "ocr_ok" in r.flags
        assert r.texto is not None
        assert "12345" in r.texto or "FACTURA" in r.texto.upper()
