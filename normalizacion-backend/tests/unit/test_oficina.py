"""Formatos de oficina que estaban MUDOS (RTF, ODT/ODS/ODP, PPTX) ahora producen texto
buscable. Existe porque un padrón en cualquiera de esos formatos iba a HOT y se indexaba
con `texto_indexable` vacío: 0 % de cobertura, indistinguible de un archivo ilegible."""

from __future__ import annotations

import io

from normalizacion.core.config import PerillasWorker
from normalizacion.ingesta.workers.extractores import ContextoExtraccion, extractor_para
from normalizacion.ingesta.workers.extractores.oficina import (
    extraer_odf,
    extraer_pptx,
    extraer_rtf,
)

CURP = "GOMC800101HDFXXX01"


def _ctx(datos: bytes, tipo: str) -> ContextoExtraccion:
    return ContextoExtraccion(
        fuente=io.BytesIO(datos), nombre="x", tipo_real=tipo, tamano=len(datos),
        perillas=PerillasWorker(),
    )


def _odt(parrafos: list[str]) -> bytes:
    from odf.opendocument import OpenDocumentText
    from odf.text import P

    d = OpenDocumentText()
    for t in parrafos:
        d.text.addElement(P(text=t))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx(textos: list[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    caja = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(5_000_000), Emu(5_000_000))
    caja.text_frame.text = "\n".join(textos)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class TestRegistro:
    def test_los_tipos_tienen_extractor(self) -> None:
        assert extractor_para("application/rtf") is extraer_rtf
        assert extractor_para("application/vnd.oasis.opendocument.text") is extraer_odf
        assert (
            extractor_para("application/vnd.openxmlformats-officedocument.presentationml.presentation")
            is extraer_pptx
        )


class TestRtf:
    def test_el_texto_del_rtf_llega(self) -> None:
        rtf = ("{\\rtf1\\ansi Registro con " + CURP + " dentro.}").encode()
        r = extraer_rtf(_ctx(rtf, "application/rtf"))
        assert r.texto and CURP in r.texto


class TestOdf:
    def test_el_texto_del_odt_llega(self) -> None:
        datos = _odt([f"linea de relleno {i}" for i in range(20)] + [f"aqui va {CURP}"])
        r = extraer_odf(_ctx(datos, "application/vnd.oasis.opendocument.text"))
        assert r.texto and CURP in r.texto


class TestPptx:
    def test_el_texto_del_pptx_llega(self) -> None:
        datos = _pptx([f"punto {i}" for i in range(5)] + [f"contacto {CURP}"])
        r = extraer_pptx(
            _ctx(datos, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        )
        assert r.texto and CURP in r.texto
