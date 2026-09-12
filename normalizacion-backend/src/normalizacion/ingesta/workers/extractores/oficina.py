"""Extractores de formatos de OFICINA que estaban MUDOS (sin plugin → `sin_extractor_l1`,
0 % de contenido buscable pese a ir a HOT): RTF, ODT/ODS/ODP (OpenDocument) y PPTX.

Ya estaban en la lista blanca (`tipos_interes`) — van a HOT — pero sin extractor se
indexaban vacíos. Producen texto plano buscable, acotado por `extractor_max_chars` (el
troceado de los grandes en N docs es un paso posterior; aquí lo importante es pasar de
mudo a cubierto). Todo degrada con flags: sin la librería, el archivo se indexa igual con
lo que haya (`extraer` devuelve `sin_extractor_l1` solo si NO hay plugin).
"""

from __future__ import annotations

from typing import Any

from . import ContextoExtraccion, ResultadoExtraccion, registrar


def _acotar(texto: str, maximo: int) -> tuple[str | None, list[str]]:
    flags = ["texto_truncado"] if len(texto) > maximo else []
    return (texto[:maximo].strip() or None), flags


@registrar("application/rtf", "text/rtf")
def extraer_rtf(ctx: ContextoExtraccion) -> ResultadoExtraccion:
    from striprtf.striprtf import rtf_to_text

    maximo = ctx.perillas.extractor_max_chars
    datos = ctx.fuente.read(maximo * 4)
    # El RTF es ASCII con escapes; latin-1 nunca falla y los \'xx se resuelven dentro.
    texto = rtf_to_text(datos.decode("latin-1", "replace"))
    doc, flags = _acotar(texto, maximo)
    return ResultadoExtraccion(campos={"lineas": (doc or "").count("\n") + 1}, texto=doc, flags=flags)


@registrar(
    "application/vnd.oasis.opendocument.text",
    "application/vnd.oasis.opendocument.spreadsheet",
    "application/vnd.oasis.opendocument.presentation",
)
def extraer_odf(ctx: ContextoExtraccion) -> ResultadoExtraccion:
    """ODT/ODS/ODP: `teletype.extractText` recorre el cuerpo y saca TODO el texto
    (párrafos, celdas, cuadros de diapositiva) de forma uniforme."""
    from odf import teletype
    from odf.opendocument import load

    documento = load(ctx.fuente)
    texto = teletype.extractText(documento.body)
    maximo = ctx.perillas.extractor_max_chars
    doc, flags = _acotar(texto, maximo)
    return ResultadoExtraccion(campos={"tipo_odf": ctx.tipo_real.rsplit(".", 1)[-1]}, texto=doc, flags=flags)


@registrar("application/vnd.openxmlformats-officedocument.presentationml.presentation")
def extraer_pptx(ctx: ContextoExtraccion) -> ResultadoExtraccion:
    from pptx import Presentation

    prs = Presentation(ctx.fuente)
    partes: list[str] = []
    for diapo in prs.slides:
        for forma in diapo.shapes:
            if forma.has_text_frame and forma.text_frame.text.strip():
                partes.append(forma.text_frame.text)
    campos: dict[str, Any] = {"diapositivas": len(prs.slides._sldIdLst)}  # noqa: SLF001
    maximo = ctx.perillas.extractor_max_chars
    doc, flags = _acotar("\n".join(partes), maximo)
    return ResultadoExtraccion(campos=campos, texto=doc, flags=flags)
